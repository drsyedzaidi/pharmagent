"""Report Agent tool: a one-page briefing memo whose every number is traced.

The memo is a template filled only from PharmState (app.compute.memo_build); a
language model never writes a number. Each numeric cell carries a small grey tag
``[#k]`` with the index of the audit entry that produced it. Before anything is
written the finished DOCX is read back and its text checked with
``memo_untraced_numbers``: a number that is not in the state-derived value table
(or a tag that names a non-existent audit entry) raises ``MemoTraceError`` and
nothing is saved or audited. The registry passes this tool a read-only snapshot
of the audit chain (``uses_audit_trail``); run directly, tags read ``[#?]``.
"""
from __future__ import annotations

import hashlib
import io
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt, RGBColor
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph

from app.compute.memo import (
    TAG_RE,
    AuditTrace,
    MemoTraceError,
    memo_bad_trace_tags,
    memo_untraced_numbers,
)
from app.compute.memo_build import DEFAULT_TITLE, Memo, build_memo
from app.compute.memo_model import Bullet, Heading, Para, Table
from app.core.pharmstate import PharmState
from app.tools.base import Tool, ToolContext, ToolResult

TABLE_STYLE = "Light Grid Accent 1"
TAG_PT = 7
TAG_GREY = RGBColor(0x80, 0x80, 0x80)
_SAFE = re.compile(r"[^A-Za-z0-9_-]")
MAX_REPORTED = 12          # untraced tokens quoted in an error message


def _add_text(paragraph: DocxParagraph, text: str, *, bold: bool = False) -> None:
    """Add ``text``; each trace tag becomes its own small grey run."""
    pos = 0
    for m in TAG_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()]).bold = bold or None
        tag = paragraph.add_run(m.group(0))
        tag.font.size, tag.font.color.rgb = Pt(TAG_PT), TAG_GREY
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:]).bold = bold or None


def render_docx(memo: Memo) -> bytes:
    doc = Document()
    doc.core_properties.title = memo.title
    doc.core_properties.author = "PharmAgent"
    for block in memo.blocks():
        if isinstance(block, Heading):
            doc.add_heading(block.text, level=block.level)
        elif isinstance(block, Para):
            _add_text(doc.add_paragraph(), block.text)
        elif isinstance(block, Bullet):
            _add_text(doc.add_paragraph(style="List Bullet"), block.text)
        elif isinstance(block, Table):
            table = doc.add_table(rows=1, cols=len(block.headers))
            table.style = TABLE_STYLE
            for cell, header in zip(table.rows[0].cells, block.headers):
                _add_text(cell.paragraphs[0], header, bold=True)
            for row in block.rows:
                for cell, text in zip(table.add_row().cells, row):
                    _add_text(cell.paragraphs[0], text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def docx_plain_text(doc: Any) -> str:
    """Body text of a python-docx Document in order: one line per paragraph, one
    ``a | b | c`` line per table row. Matches ``Memo.text()`` for a rendered memo."""
    lines: list[str] = []
    for item in doc.iter_inner_content():
        if isinstance(item, DocxTable):
            lines += [" | ".join(c.text for c in row.cells) for row in item.rows]
        elif isinstance(item, DocxParagraph):
            lines.append(item.text)
    return "\n".join(lines)


def _filename(state: PharmState) -> str:
    meta = state.dataset_metadata or {}
    parts = [_SAFE.sub("", str(p))[:40] for p in (state.session_id, meta.get("dataset_id")) if p]
    return "memo_" + ("_".join(p for p in parts if p) or "analysis") + ".docx"


def _verify(memo: Memo, data: bytes, n_entries_after: int) -> str:
    """Read the rendered DOCX back and fail closed on any untraced number or bad tag."""
    text = docx_plain_text(Document(io.BytesIO(data)))
    untraced = memo_untraced_numbers(text, memo.values)
    bad = memo_bad_trace_tags(text, n_entries_after)
    if untraced or bad:
        shown = ", ".join(untraced[:MAX_REPORTED]) + (" ..." if len(untraced) > MAX_REPORTED else "")
        raise MemoTraceError(
            "briefing memo refused: "
            + (f"{len(untraced)} number(s) not traceable to analysis state ({shown}); " if untraced else "")
            + (f"trace tag(s) naming no audit entry ({', '.join(bad)}); " if bad else "")
            + "nothing was saved")
    return text


def build_briefing_memo(state: PharmState, ctx: ToolContext, args: dict[str, Any]) -> ToolResult:
    trace = AuditTrace(ctx.audit_trail)
    # The registry appends this call's own entry right after run(): its index is the
    # current length. An empty snapshot means a direct call outside the registry.
    own = trace.n_entries if trace.n_entries else None
    memo = build_memo(state, trace, title=args.get("title") or DEFAULT_TITLE, own_index=own)
    if not memo.sections:
        raise ValueError("nothing to brief: run at least one analysis (dataset, NCA, QC, model fit) first")
    data = render_docx(memo)
    text = _verify(memo, data, trace.n_entries + 1)

    out_dir = Path(ctx.data_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / _filename(state)
    path.write_bytes(data)

    result = {
        "status": "ok",
        "memo_path": str(path),
        "filename": path.name,
        "title": memo.title,
        "sections_included": [s.key for s in memo.sections],
        "sections_omitted": [{"key": k, "title": t, "note": "not run"} for k, t in memo.omitted],
        "results_withheld": [{"key": k, "title": t, "note": "computed on a previously loaded dataset"}
                             for k, t in memo.withheld],
        "dataset_id": state.dataset_id,
        "n_values": sum(1 for v in memo.values if v.kind == "number"),
        "n_inputs": sum(1 for v in memo.values if v.kind == "input"),
        "n_literals": sum(1 for v in memo.values if v.kind == "literal"),
        "audit_index": own,
        "sha256": hashlib.sha256(data).hexdigest(),
        "untraced_numbers": [],
        "bad_tags": [],
        "memo_text": text,
        "value_table": [{"key": v.key, "text": v.text, "audit_index": v.audit_index, "kind": v.kind}
                        for v in memo.values],
    }
    return ToolResult(
        summary=(f"Built briefing memo at {path}: {len(memo.sections)} section(s), "
                 f"{len(memo.omitted)} not run"
                 + (f", {len(memo.withheld)} result(s) from a previously loaded dataset withheld"
                    if memo.withheld else "")
                 + "; every number traced to analysis state."),
        action=f"build_briefing_memo -> {path}",
        writes={"memo_path": str(path), "memo_results": result},
        result=result,
    )


TOOLS = [
    Tool("build_briefing_memo",
         "Build a briefing memo (DOCX) from the current analysis state: dataset, NCA "
         "(geometric mean CL/F and AUC), QC verdict with failing checks, structural "
         "model comparison, population fit and covariate effects, diagnostics headline "
         "numbers, indirect comparisons, exposure-response fits and the dose-selection "
         "decision. A fixed template filled only from analysis state; sections that were "
         "not run are omitted with a note, and results computed on a previously loaded "
         "dataset are withheld. User-supplied inputs are marked as such. Every number "
         "carries the audit-entry tag that produced it; the tool refuses to build if any "
         "number cannot be traced. Never writes a number itself.",
         "report",
         {"type": "object",
          "properties": {"title": {"type": "string", "description": "memo title (no digits)"}},
          "required": []},
         build_briefing_memo, uses_audit_trail=True),
]
