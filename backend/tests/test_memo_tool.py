"""build_briefing_memo: registry tool, DOCX output, fail-closed trace check, and the
end-to-end build on theoph_pk.csv through the orchestrator with no LLM key."""
import itertools
import re
from pathlib import Path

import pytest
from docx import Document

import app.compute.memo_build as memo_build
from app.compute.memo import TAG_RE, MemoTraceError, TracedValue, memo_untraced_numbers
from app.compute.memo_build import SECTION_TITLES
from app.compute.memo_model import Para, Section
from app.core.llm import MockLLM
from app.core.orchestrator import Orchestrator
from app.core.pharmstate import AGENT_WRITE_FIELDS, PharmState
from app.core.store import SessionStore
from app.tools.base import ToolContext
from app.tools.builtins import default_registry
from app.tools.memo_tools import docx_plain_text
from tests.memo_fixtures import IDX, full_chain, full_state

THEOPH = str(Path(__file__).parent.parent / "sample_data" / "theoph_pk.csv")


def _run(state, tmp_path, args=None, chain=None):
    chain = chain if chain is not None else full_chain()
    new_state, res = default_registry().execute(
        "build_briefing_memo", state=state, ctx=ToolContext(data_dir=str(tmp_path)),
        args=args or {}, audit=chain, timestamp="t-memo", actor="tester")
    return new_state, res, chain


# ── registration ───────────────────────────────────────────────────────────────
def test_registered_cheap_report_tool_that_cites_the_audit_trail():
    tool = default_registry().get("build_briefing_memo")
    assert tool.agent == "report"
    assert tool.expensive is False and tool.proposable is False
    assert tool.uses_audit_trail is True
    assert {"memo_path", "memo_results"} <= AGENT_WRITE_FIELDS["report"]
    assert set(tool.input_schema["properties"]) == {"title"}


# ── happy path on a full synthetic state ───────────────────────────────────────
def test_builds_a_docx_and_records_state_and_a_matching_audit_entry(tmp_path):
    state, res, chain = _run(full_state(), tmp_path)
    r = res.result
    path = Path(r["memo_path"])
    assert path.exists() and path.suffix == ".docx" and path.parent == tmp_path
    assert state.memo_path == str(path) and state.memo_results == r
    assert r["status"] == "ok" and r["untraced_numbers"] == [] and r["bad_tags"] == []
    assert r["sections_included"] == list(SECTION_TITLES) and r["sections_omitted"] == []
    # the memo names its own audit entry, and that is exactly the entry the registry appended
    entry = chain.entries[-1]
    assert (entry.tool, entry.agent, entry.actor) == ("build_briefing_memo", "report", "tester")
    assert r["audit_index"] == entry.index == len(chain.entries) - 1
    assert f"[#{entry.index}]" in r["memo_text"]
    assert chain.verify()


def test_docx_text_equals_stored_memo_text_and_passes_the_trace_check(tmp_path):
    _, res, _ = _run(full_state(), tmp_path)
    doc = Document(res.result["memo_path"])
    text = docx_plain_text(doc)
    assert text == res.result["memo_text"]
    table = [(v["text"], v["kind"]) for v in res.result["value_table"]]
    values = [TracedValue("k", t, None, kind) for t, kind in table]
    assert memo_untraced_numbers(text, values) == []
    assert len(doc.tables) >= 7


def test_numeric_cells_carry_a_small_grey_tag_run(tmp_path):
    _, res, _ = _run(full_state(), tmp_path)
    doc = Document(res.result["memo_path"])
    cell = next(c for t in doc.tables for row in t.rows for c in row.cells
                if re.fullmatch(r"2\.71 \[#3\]", c.text))
    runs = cell.paragraphs[0].runs
    assert [r.text for r in runs][-1] == "[#3]" and runs[-1].font.size.pt == 7
    assert runs[0].font.size is None            # the value itself is body size


def test_trace_tags_in_the_document_match_the_chain(tmp_path):
    _, res, chain = _run(full_state(), tmp_path)
    tags = {int(m.group(1)) for m in re.finditer(r"\[#(\d+)\]", res.result["memo_text"])}
    assert tags == {IDX["load"], IDX["profile"], IDX["nca"], IDX["qc"], IDX["pk"], IDX["nlme"],
                    IDX["scm"], IDX["forest"], IDX["vpc"], IDX["diag"], IDX["ind_hr"],
                    IDX["ind_md"], IDX["er_eff"], IDX["er_tox"], IDX["er_pfs"], IDX["er_boot"],
                    IDX["dose"], len(chain.entries) - 1}
    assert max(tags) < len(chain.entries)


def test_custom_title_and_digit_titles_are_rejected(tmp_path):
    _, res, _ = _run(full_state(), tmp_path, {"title": "Exposure briefing"})
    assert res.result["title"] == "Exposure briefing"
    assert docx_plain_text(Document(res.result["memo_path"])).startswith("Exposure briefing")
    with pytest.raises(ValueError, match="digits"):
        _run(full_state(), tmp_path, {"title": "Study 101"})


def test_missing_sections_are_omitted_with_a_one_line_note(tmp_path):
    state = PharmState(qc_verdict="PASS", nca_summary=full_state().nca_summary)
    _, res, _ = _run(state, tmp_path)
    assert res.result["sections_included"] == ["nca", "qc"]
    omitted = {o["key"]: o["title"] for o in res.result["sections_omitted"]}
    assert set(omitted) == set(SECTION_TITLES) - {"nca", "qc"}
    assert "Population fit: not run." in res.result["memo_text"]


def test_nothing_to_brief_raises_before_writing_or_auditing(tmp_path):
    chain = full_chain()
    n = len(chain.entries)
    with pytest.raises(ValueError, match="nothing to brief"):
        _run(PharmState(), tmp_path, chain=chain)
    assert len(chain.entries) == n and list(tmp_path.iterdir()) == []


def test_filename_is_safe_and_deterministic(tmp_path):
    state = full_state().model_copy(update={"session_id": "../evil/s1"})
    _, res, _ = _run(state, tmp_path)
    name = Path(res.result["memo_path"]).name
    assert re.fullmatch(r"memo_[A-Za-z0-9_-]+\.docx", name) and ".." not in name
    assert res.result["filename"] == name


# ── fail closed ────────────────────────────────────────────────────────────────
def test_an_untraced_number_in_the_memo_raises_and_leaves_no_file_or_audit_entry(tmp_path, monkeypatch):
    def rogue(state, trace, rec):
        return Section("rogue", "Rogue", (Para("The AUC was 99.9 in all subjects."),), "")

    monkeypatch.setattr(memo_build, "BUILDERS", (*memo_build.BUILDERS, ("rogue", rogue)))
    monkeypatch.setitem(memo_build.SECTION_TITLES, "rogue", "Rogue")
    chain = full_chain()
    n = len(chain.entries)
    with pytest.raises(MemoTraceError, match=r"99\.9"):
        _run(full_state(), tmp_path, chain=chain)
    assert len(chain.entries) == n and list(tmp_path.iterdir()) == []


def test_a_tag_pointing_beyond_the_chain_raises(tmp_path, monkeypatch):
    def rogue(state, trace, rec):
        return Section("rogue", "Rogue", (Para("Value tagged [#999]."),), "")

    monkeypatch.setattr(memo_build, "BUILDERS", (*memo_build.BUILDERS, ("rogue", rogue)))
    monkeypatch.setitem(memo_build.SECTION_TITLES, "rogue", "Rogue")
    with pytest.raises(MemoTraceError, match=r"\[#999\]"):
        _run(full_state(), tmp_path)


# ── end to end through the orchestrator, keyless, on theoph_pk.csv ──────────────
def _orch():
    counter = itertools.count()
    return Orchestrator(llm=MockLLM(), clock=lambda: f"t{next(counter)}",
                        store=SessionStore(":memory:"))


def test_memo_builds_end_to_end_on_theoph_through_the_orchestrator(tmp_path, monkeypatch):
    from app.config import settings
    monkeypatch.setattr(settings, "data_dir", tmp_path)
    orch = _orch()
    sid = orch.create_session().id
    orch.start_workflow(sid, "nca_full", {"path": THEOPH})           # load -> ... -> QC gate
    orch.run_tool(sid, "fit_pk_model", "modeler", {"compare": True})
    orch.run_tool(sid, "run_vpc", "modeler", {})
    orch.run_tool(sid, "run_diagnostics", "modeler", {})
    orch.run_tool(sid, "indirect_comparison", "statistician", {
        "ab": {"estimate": 0.5, "se": 0.2}, "cb": {"estimate": 0.8, "se": 0.3}, "scale": "HR"})

    out = orch.run_tool(sid, "build_briefing_memo", "report", {})
    res, sess = out["result"], orch.get_session(sid)

    assert out["audit_ok"] is True and res["status"] == "ok" and res["untraced_numbers"] == []
    assert Path(res["memo_path"]).exists()
    assert res["sections_included"] == ["dataset", "nca", "qc", "structural", "diagnostics", "indirect"]
    assert {o["key"] for o in res["sections_omitted"]} == {"population", "covariates", "er", "dose_selection"}
    assert out["state"]["memo_path"] == res["memo_path"]

    # every tag resolves to a real audit entry, and the tool that produced it is the right one
    entries = {e.index: e for e in sess.audit.entries}
    text = res["memo_text"]
    assert TAG_RE.search(text) and "[#?]" not in text
    for line, tool in (("CL/F |", "compute_nca"), ("Subjects |", "load_dataset"),
                       ("QC verdict", "run_qc"), ("Best model by AIC", "fit_pk_model"),
                       ("R² of IPRED", "run_vpc"), ("IWRES mean", "run_diagnostics"),
                       ("Indirect |", "indirect_comparison")):
        row = next(ln for ln in text.splitlines() if ln.startswith(line))
        idx = {int(i) for i in re.findall(r"\[#(\d+)\]", row)}
        assert idx and all(entries[i].tool == tool for i in idx), (line, idx)
    assert entries[res["audit_index"]].tool == "build_briefing_memo"

    # hand-check against the NCA the memo was built from
    cl_f = next(g for g in sess.state.nca_summary["descriptive"] if g["group"] == "all")
    gm = next(p for p in cl_f["parameters"] if p["parameter"] == "CL_F")["geomean"]
    assert f"{gm:.2f} [#" in text
    assert memo_untraced_numbers(text, _values(res)) == []


def _values(res):
    return [TracedValue(v["key"], v["text"], v["audit_index"], v["kind"]) for v in res["value_table"]]


def test_agent_state_summary_tells_the_llm_a_memo_exists(tmp_path):
    from app.agents.definitions import AGENTS
    state, _, _ = _run(full_state(), tmp_path)
    assert AGENTS["report"]._state_summary(state)["memo_path"] == state.memo_path
    assert AGENTS["report"]._state_summary(PharmState())["memo_path"] is None


def test_result_carries_the_sha256_of_the_saved_file(tmp_path):
    import hashlib
    _, res, _ = _run(full_state(), tmp_path)
    assert res.result["sha256"] == hashlib.sha256(Path(res.result["memo_path"]).read_bytes()).hexdigest()


def test_unrenderable_state_is_a_value_error_not_an_attribute_error(tmp_path):
    with pytest.raises(ValueError, match="cannot render the"):
        _run(PharmState(dataset_metadata={"n_subjects": 3, "dose_levels": 5}), tmp_path)
