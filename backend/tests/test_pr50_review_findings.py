"""Tests for the post-merge review findings fixed in PR #50.

Covers:
  Item 1 (HIGH): report disposition uses audit trail, not "blocks gate"
    - Adversarial-review gate approved → renders actor, timestamp, reason
    - No gate approval recorded → renders "no gate decision recorded"
    - Earlier gate (plan/special-pop) approved but adversarial gate pending →
      does NOT pick up the wrong approval
    - Adversarial gate rejected → renders "no gate decision recorded"
    - The string "blocks gate" never appears in rendered output
  Item 3 (MEDIUM): run_dose_sweep from_protocol_plan flag
    - Default (no flag) with protocol_sim populated → same doses/writes as
      protocol_sim absent
    - With flag → plan doses are used and protocol_sim is written
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document

from app.core.pharmstate import PharmState
from app.tools.base import ToolContext
from app.tools.report_tools import _find_gate_approval, _protocol_sim_section

# ── Helpers ──────────────────────────────────────────────────────────────────

def _adversarial_review_state() -> PharmState:
    """State with an unresolved HIGH adversarial finding."""
    return PharmState(
        protocol_sim={
            "source_fit": {"fit_id": "fit_01", "label": "1cmt", "theta": {}},
            "plan": {"bands": [], "regimens": [], "n_subjects": 50,
                     "replicates": 10, "seed": 42, "uncertainty_method": "bootstrap"},
            "pta": {"rows": [], "n_subjects": 50, "replicates": 10, "seed": 42,
                    "recommendations": []},
            "provenance": {"source_fit_id": "fit_01", "review_hash": "abc",
                           "plan_hash": "def", "seed": 42,
                           "tool_versions": {}, "wall_clock": {}},
        },
        review_results={
            "goal": "adversarial",
            "goal_met": False,
            "findings": [
                {"id": "F1", "severity": "HIGH", "target": "pta_mc_se",
                 "evidence": "MC SE too large", "resolved": False},
            ],
            "counts": {"CRITICAL": 0, "HIGH": 1, "MEDIUM": 0, "LOW": 0},
        },
    )


def _audit_entry(action: str, label: str, actor: str = "dr.z",
                 timestamp: str = "2026-10-25T12:00:00Z",
                 reason: str = "reviewed and accepted") -> dict[str, Any]:
    return {
        "agent": "qc",
        "tool": "human_review",
        "action": action,
        "inputs": {"approve": action.startswith("approved"),
                   "review": {"after_step": 6, "label": label}},
        "outputs": {"workflow": "protocol_simulation"},
        "actor": actor,
        "timestamp": timestamp,
        "reason": reason,
    }


# ── Item 1: report disposition tests ────────────────────────────────────────

class TestGateApprovalLookup:
    """Tests for _find_gate_approval matching the adversarial-review gate."""

    def test_adversarial_gate_approved(self):
        """An approval whose label contains 'Adversarial review' is found."""
        trail = (
            _audit_entry("approved after step 1",
                         "Propose simulation plan (human review gate)",
                         actor="alice", timestamp="t1", reason="plan ok"),
            _audit_entry("approved after step 6",
                         "Adversarial review of simulation claims",
                         actor="dr.z", timestamp="t2", reason="findings acceptable"),
        )
        result = _find_gate_approval(trail)
        assert result is not None
        assert result["actor"] == "dr.z"
        assert result["timestamp"] == "t2"
        assert result["reason"] == "findings acceptable"

    def test_no_adversarial_approval_returns_none(self):
        """When the adversarial gate was never approved, return None."""
        trail = (
            _audit_entry("approved after step 1",
                         "Propose simulation plan (human review gate)",
                         actor="alice"),
        )
        result = _find_gate_approval(trail)
        assert result is None

    def test_adversarial_rejected_returns_none(self):
        """A rejection of the adversarial gate is not treated as an approval."""
        trail = (
            _audit_entry("rejected after step 6",
                         "Adversarial review of simulation claims",
                         actor="dr.z", reason="needs rework"),
        )
        result = _find_gate_approval(trail)
        assert result is None

    def test_earlier_gate_not_confused_with_adversarial(self):
        """If only an earlier gate (special-pop) is approved, don't return it."""
        trail = (
            _audit_entry("approved after step 4",
                         "Special-population / pediatric exposures",
                         actor="bob", timestamp="t1", reason="sp ok"),
        )
        result = _find_gate_approval(trail)
        assert result is None

    def test_empty_trail(self):
        result = _find_gate_approval(())
        assert result is None


class TestReportDisposition:
    """Tests for the rendered finding disposition in _protocol_sim_section."""

    def test_approved_renders_actor_timestamp_reason(self):
        state = _adversarial_review_state()
        trail = (
            _audit_entry("approved after step 6",
                         "Adversarial review of simulation claims",
                         actor="dr.zaidi", timestamp="2026-10-25T15:00:00Z",
                         reason="risk accepted per protocol"),
        )
        result = _protocol_sim_section(Document(), state, 5, audit_trail=trail)
        md = result.get("rendered_markdown", "")
        assert "dr.zaidi" in md
        assert "2026-10-25T15:00:00Z" in md
        assert "risk accepted per protocol" in md
        assert "blocks gate" not in md.lower()

    def test_no_approval_renders_no_gate_decision(self):
        state = _adversarial_review_state()
        trail = ()  # empty audit trail
        result = _protocol_sim_section(Document(), state, 5, audit_trail=trail)
        md = result.get("rendered_markdown", "")
        assert "no gate decision recorded" in md
        assert "blocks gate" not in md.lower()

    def test_blocks_gate_never_appears(self):
        """Even with an earlier non-adversarial approval, 'blocks gate' is gone."""
        state = _adversarial_review_state()
        trail = (
            _audit_entry("approved after step 1",
                         "Propose simulation plan (human review gate)",
                         actor="alice"),
        )
        result = _protocol_sim_section(Document(), state, 5, audit_trail=trail)
        md = result.get("rendered_markdown", "")
        assert "blocks gate" not in md.lower()
        assert "no gate decision recorded" in md

    def test_resolved_finding_still_shows_resolved(self):
        state = _adversarial_review_state()
        state.review_results["findings"][0]["resolved"] = True
        trail = ()
        result = _protocol_sim_section(Document(), state, 5, audit_trail=trail)
        md = result.get("rendered_markdown", "")
        assert "Resolved" in md
        assert "no gate decision" not in md


# ── Item 3: dose-sweep from_protocol_plan tests ─────────────────────────────

def _nlme_state_for_sweep() -> PharmState:
    """State with a fitted model and a protocol_sim plan that sets custom doses."""
    return PharmState(
        nlme_results={
            "status": "ok", "converged": True,
            "fit_id": "fit_sweep_01", "model_key": "oral_1cmt",
            "theta": {"CL": 5.0, "V": 50.0, "KA": 1.0},
            "omega_cv_pct": {"CL": 20.0, "V": 15.0},
            "iiv_params": ["CL", "V"],
            "sigma": {"prop": 0.1, "add": 0.0},
        },
        pk_model_results={
            "status": "ok",
            "mode": "fit",
            "model_key": "oral_1cmt",
            "population": {
                "parameters": {
                    "CL": {"typical_value": 5.0},
                    "V": {"typical_value": 50.0},
                    "KA": {"typical_value": 1.0},
                }
            },
        },
        protocol_sim={
            "source_fit": {"fit_id": "fit_sweep_01", "model_key": "oral_1cmt",
                           "theta": {"CL": 5.0, "V": 50.0, "KA": 1.0}},
            "plan": {
                "regimens": [
                    {"dose_mg": 75.0, "tau_h": 12.0, "n_doses": 2},
                    {"dose_mg": 150.0, "tau_h": 12.0, "n_doses": 2},
                    {"dose_mg": 300.0, "tau_h": 12.0, "n_doses": 2},
                ],
                "bands": [], "n_subjects": 50, "replicates": 10, "seed": 42,
            },
            "provenance": {"tool_versions": {}, "wall_clock": {}},
        },
    )


class TestDoseSweepFlag:
    """from_protocol_plan flag isolates ordinary sweeps from plan data."""

    def test_default_no_flag_ignores_plan(self, tmp_path: Path):
        """Without from_protocol_plan, protocol_sim is NOT read even if present."""
        from app.tools.pkmodel_tools import run_dose_sweep

        state = _nlme_state_for_sweep()
        ctx = ToolContext(data_dir=str(tmp_path))
        # Default sweep (no flag) — should use default doses, not plan doses
        result = run_dose_sweep(state, ctx, {"dose": 100.0})
        doses_used = result.result.get("doses", [])

        # Should NOT have the plan doses [75, 150, 300]
        assert 75.0 not in doses_used
        # Should NOT write protocol_sim
        assert "protocol_sim" not in result.writes

    def test_flag_true_reads_plan(self, tmp_path: Path):
        """With from_protocol_plan=True, plan doses are used and protocol_sim is written."""
        from app.tools.pkmodel_tools import run_dose_sweep

        state = _nlme_state_for_sweep()
        ctx = ToolContext(data_dir=str(tmp_path))
        result = run_dose_sweep(state, ctx, {"from_protocol_plan": True})
        doses_used = result.result.get("doses", [])

        # Plan has regimens at 75, 150, 300
        assert sorted(doses_used) == [75.0, 150.0, 300.0]
        # protocol_sim must be written
        assert "protocol_sim" in result.writes
        ps = result.writes["protocol_sim"]
        assert "dose_sweep" in ps
        assert ps["provenance"]["tool_versions"]["run_dose_sweep"] == "1.0.0"

    def test_default_matches_absent_protocol_sim(self, tmp_path: Path):
        """Default path output is identical whether protocol_sim is set or not."""
        from app.tools.pkmodel_tools import run_dose_sweep

        state_with = _nlme_state_for_sweep()
        state_without = state_with.model_copy(update={"protocol_sim": None})
        ctx = ToolContext(data_dir=str(tmp_path))
        args = {"dose": 100.0}

        res_with = run_dose_sweep(state_with, ctx, dict(args))
        res_without = run_dose_sweep(state_without, ctx, dict(args))

        assert res_with.result["doses"] == res_without.result["doses"]
        assert "protocol_sim" not in res_with.writes
        assert "protocol_sim" not in res_without.writes
