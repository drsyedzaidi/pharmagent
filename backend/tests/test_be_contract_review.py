"""Independent review of 3daebfe (BE safety contract): every finding reproduced here first, then fixed.

B1 replicate rows merged silently without a period column; B2 the estimator rounds to 4 dp before the contract rounds;
H1 rows dropped in the tool with no ledger entry; H2 malformed arguments escape as exceptions; H3 near-default custom
limits pass as the default; M1 bools counted as values; M2 ambiguous level text; M3 reports show the unrounded CI
beside a verdict taken from the rounded one; L1 duplicate pre-specified exclusions."""
from __future__ import annotations

import math

import numpy as np
import pandas as pd
import pytest

import app.compute.bioequivalence as est
from app.compute.be_contract import BEContractError, assess_average_be
from tests.test_be_tool_contract import _call, _frame

FULL = {"design": "crossover", "profile": "fda_2026", "test": "B", "reference": "A", "period_col": "PERIOD"}


def _pframe(n=12, **kw):
    return _frame(n, **kw)                                    # two-sequence crossover with a PERIOD column


def _recs(n=12):
    out = []
    for i in range(n):
        out += [{"subject": i, "treatment": "R", "params": {"Cmax": 100.0 + i}},
                {"subject": i, "treatment": "T", "params": {"Cmax": 101.0 + i}}]
    return out


def _contract(recs, **kw):
    return assess_average_be(recs, **{"design": "crossover", "profile": "fda_2026", "test": "T", "reference": "R",
                                      **kw})


# ── B1 ─────────────────────────────────────────────────────────────────────────────────────────
def test_a_crossover_needs_a_period_column():
    res = _call(_pframe(), **{k: v for k, v in FULL.items() if k != "period_col"}).result
    assert res["status"] == "missing_required_args" and res["missing"] == ["period_col"]


def test_duplicated_sampling_rows_are_refused_not_pooled():
    df = _pframe(13)
    dup = df[(df.ID == 1) & (df.TRT == "A")].assign(DV=lambda d: d.DV * 3)
    res = _call(pd.concat([df, dup]), **FULL).result
    assert res["status"] == "unsupported_design" and "duplicate" in res["message"]


def test_test_and_reference_in_the_same_period_is_not_a_crossover():
    df = _pframe().assign(PERIOD=1)
    res = _call(df, **FULL).result
    assert res["status"] == "invalid_arguments" and "period" in res["message"]


# ── B2 ─────────────────────────────────────────────────────────────────────────────────────────
def test_the_estimator_exposes_unrounded_bounds():
    r = est.be_one_parameter([100.0, 103.0, 98.0], [101.0, 99.0, 100.0], paired=True)
    assert round(r["ci_lower_raw"], 4) == r["ci_lower_pct"]


def test_the_verdict_is_not_double_rounded(monkeypatch):
    # a true lower bound of 79.99496 is 79.99 at two decimals; a 4-dp copy (79.995) would round to 80.00
    monkeypatch.setattr(est, "be_one_parameter", lambda *a, **k: {
        "gmr_pct": 95.0, "ci_lower_pct": 79.995, "ci_upper_pct": 110.0, "cv_intra_pct": 10.0,
        "gmr_raw": 95.0, "ci_lower_raw": 79.99496, "ci_upper_raw": 110.0})
    import app.compute.be_contract as bc
    monkeypatch.setattr(bc, "be_one_parameter", est.be_one_parameter)
    p = _contract(_recs())["parameters"]["Cmax"]
    assert p["ci_lower_rounded"] == "79.99" and p["within_limits"] is False
    m13a = _contract(_recs(), profile="ich_m13a")["parameters"]["Cmax"]
    assert m13a["within_limits"] is False


# ── H1 ─────────────────────────────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("raw, reason", [("BQL", "non_numeric_dv"), (float("nan"), "missing_dv"),
                                         (float("inf"), "non_finite_dv")])
def test_every_dropped_concentration_row_is_in_the_ledger(raw, reason):
    df = _pframe(13).astype({"DV": object})
    idx = df[(df.ID == 3) & (df.TRT == "B") & (df.TIME == 2.0)].index[0]
    df.loc[idx, "DV"] = raw
    res = _call(df, **FULL).result
    assert res["status"] == "ok"
    hits = [e for e in res["exclusions"] if e["subject"] == 3 and e["treatment"] == "B" and e["reason"] == reason]
    assert len(hits) == 1 and hits[0]["time"] == 2.0
    assert hits[0]["original_value"] == raw or (isinstance(raw, float) and math.isnan(raw))


def test_a_row_without_subject_or_treatment_is_refused():
    df = _pframe(13).astype({"TRT": object})
    df.loc[df.index[0], "TRT"] = None
    res = _call(df, **FULL).result
    assert res["status"] == "invalid_arguments" and "treatment" in res["message"]


# ── H2 ─────────────────────────────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("extra", [
    {"limits": [1, 2, 3]}, {"limits": [float("nan"), 125]}, {"limits": 80}, {"limits": ["a", "b"]},
    {"prespecified_exclusions": ["1"]}, {"prespecified_exclusions": {"subject": 1}},
    {"prespecified_exclusions": [{"subject": [1], "reason": "x"}]},
])
def test_malformed_arguments_are_refusals_not_exceptions(extra):
    res = _call(_pframe(), **FULL, **extra).result
    assert res["status"] == "invalid_arguments"


@pytest.mark.parametrize("value, reason", [(float("inf"), "non_finite_value"), ("12", "invalid_type"),
                                           (True, "invalid_type"), (np.bool_(True), "invalid_type")])
def test_contract_values_of_the_wrong_kind_are_ledgered(value, reason):
    recs = _recs(13)
    recs[0]["params"]["Cmax"] = value
    out = _contract(recs)
    assert any(e["subject"] == 0 and e["treatment"] == "R" and e["reason"] == reason for e in out["exclusions"])
    assert out["parameters"]["Cmax"]["n_evaluable"] == 12


# ── H3 ─────────────────────────────────────────────────────────────────────────────────────────
def test_near_default_limits_are_custom_limits():
    res = _call(_pframe(), **FULL, limits=[79.996, 125.004]).result
    assert res["status"] == "invalid_arguments"
    res = _call(_pframe(), **FULL, limits=[80, 125]).result
    assert res["status"] == "ok" and res["limits"] == ["80.00", "125.00"]


# ── M2 / L1 ────────────────────────────────────────────────────────────────────────────────────
def test_level_text_matching_two_data_levels_is_refused():
    df = _pframe().astype({"TRT": object})
    df.loc[df.TRT == "B", "TRT"] = 1
    df.loc[(df.ID == 1) & (df.TRT == 1), "TRT"] = "1"
    res = _call(df, **{**FULL, "test": "1"}).result
    assert res["status"] == "invalid_arguments" and "ambiguous" in res["message"]


def test_duplicate_prespecified_exclusions_are_refused():
    with pytest.raises(BEContractError, match="more than once"):
        _contract(_recs(13), prespecified_exclusions=[{"subject": 1, "reason": "a"}, {"subject": 1, "reason": "b"}])


# ── M3 ─────────────────────────────────────────────────────────────────────────────────────────
def test_the_csv_export_carries_the_decision_fields():
    from app.core.exporters import _be
    from app.core.pharmstate import PharmState

    res = _call(_pframe(), **FULL).result
    csv = _be(PharmState(be_results=res))
    header = csv.splitlines()[0]
    for col in ("ci_lower_rounded", "ci_upper_rounded", "acceptance", "n_evaluable", "profile"):
        assert col in header


@pytest.mark.parametrize("which", ["report", "272"])
def test_reports_show_profile_rounded_ci_interim_model_and_exclusions(tmp_path, which):
    from docx import Document

    from app.core.pharmstate import PharmState
    from app.tools.base import ToolContext
    from app.tools.regulatory_tools import generate_272
    from app.tools.report_tools import generate_report

    res = _call(_pframe(13, drop={(0, "A")}), **FULL).result
    fn = generate_report if which == "report" else generate_272
    out = fn(PharmState(be_results=res), ToolContext(data_dir=str(tmp_path)), {})
    path = out.result.get("report_path") or out.result.get("regulatory_report_path") or out.writes.get(
        "regulatory_report_path")
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(
        c.text for t in doc.tables for row in t.rows for c in row.cells)
    cmax = res["parameters"]["Cmax"]
    assert "fda_2026" in text and "interim" in text and f"{len(res['exclusions'])} excluded value" in text
    assert f"{cmax['ci_lower_rounded']}–{cmax['ci_upper_rounded']}" in text


# ── re-review of 9e7582d ───────────────────────────────────────────────────────────────────────
def _replicate_with_nan_period(df):
    extra = df[(df.ID <= 3) & (df.TRT == "B")].assign(TIME=lambda d: d.TIME + 0.25, PERIOD=np.nan,
                                                      DV=lambda d: d.DV * 5)
    return pd.concat([df, extra])


def test_rows_with_a_missing_period_are_refused():
    res = _call(_replicate_with_nan_period(_pframe(13)), **FULL).result
    assert res["status"] == "invalid_arguments" and "period" in res["message"]


@pytest.mark.parametrize("col", ["TRT", "ID", "TIME", "DV"])
def test_period_col_cannot_be_a_role_column(col):
    res = _call(_pframe(), **{**FULL, "period_col": col}).result
    assert res["status"] == "invalid_arguments" and "period_col" in res["message"]


def test_a_single_sequence_is_not_a_randomised_crossover():
    df = _pframe().pipe(lambda d: d.assign(PERIOD=np.where(d.TRT == "A", 1, 2)))   # everyone: A then B
    res = _call(df, **FULL).result
    assert res["status"] == "invalid_arguments" and "sequence" in res["message"]


def test_two_missing_time_rows_are_ledgered_not_called_duplicates():
    df = _pframe(13)
    rows = df[(df.ID == 2) & (df.TRT == "B")].head(2).assign(TIME=np.nan)
    df = pd.concat([df.drop(rows.index), rows])
    res = _call(df, **FULL).result
    assert res["status"] == "ok"
    assert sum(e["reason"] == "missing_time" and e["subject"] == 2 for e in res["exclusions"]) == 2


def test_ledger_original_values_are_strict_json():
    import json

    df = _pframe(13)
    df.loc[df.index[3], "DV"] = np.nan
    res = _call(df, **FULL).result
    json.dumps(res, allow_nan=False)


def test_a_non_unique_index_does_not_crash():
    df = _pframe(13)
    df.loc[df.index[3], "DV"] = np.nan
    df.index = [0] * len(df)
    assert _call(df, **FULL).result["status"] == "ok"


def test_an_unrounded_profile_report_does_not_show_a_misleading_ci():
    from app.compute.be_contract import report_ci

    v = {"ci_lower_pct": 80.0, "ci_upper_pct": 110.0, "ci_lower_raw": 79.99996, "ci_upper_raw": 110.0,
         "ci_lower_rounded": None, "ci_upper_rounded": None}
    assert "79.99996" in report_ci(v)



def test_a_replicate_with_different_sampling_times_is_refused_as_replicate():
    """The same treatment in two periods, sampled at different times so the duplicate-time check does not fire:
    the period check itself must refuse it (owner gate on be_tools.py:104)."""
    df = _pframe(13)
    again = df[(df.ID == 1) & (df.TRT == "B")].assign(PERIOD=3, TIME=lambda d: d.TIME + 0.25)
    both = pd.concat([df, again])
    assert not both.duplicated(["ID", "TRT", "TIME"]).any()           # precondition: no duplicate sampling times
    res = _call(both, **FULL).result
    assert res["status"] == "unsupported_design" and "replicate" in res["message"]
    assert "bioequivalent" not in res and "No bioequivalence conclusion was calculated" in res["message"]
