"""Reports and exports carry the roles, QC findings and the subject-level disposition table."""
from __future__ import annotations

import csv
import io

from docx import Document

from app.core import exporters
from app.core.pharmstate import PharmState
from app.tools.base import ToolContext
from app.tools.regulatory_tools import generate_272
from app.tools.report_tools import generate_report
from tests.test_be_qc_tool import _call, _cmax, _set_dv
from tests.test_be_tool_contract import FULL, _frame


def _result_with_an_exclusion():
    df = _frame(13)
    df = _set_dv(df, 3, "B", 0.0, 0.2 * _cmax(df, 3, "B"))
    return _call(df, **FULL).result


def _doc_text(path):
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
        c.text for t in doc.tables for row in t.rows for c in row.cells)


def _report_text(fn, res, tmp_path):
    out = fn(PharmState(be_results=res), ToolContext(data_dir=str(tmp_path)), {})
    path = out.result.get("report_path") or out.result.get("regulatory_report_path") or out.writes.get(
        "regulatory_report_path")
    return _doc_text(path)


def test_both_reports_show_the_disposition_table_and_the_qc_finding(tmp_path):
    res = _result_with_an_exclusion()
    for fn in (generate_report, generate_272):
        text = _report_text(fn, res, tmp_path)
        assert "Subject disposition" in text and "predose_gt_5pct_cmax" in text, fn.__name__
        assert "excluded" in text and "decision" in text and "supportive" in text, fn.__name__


def test_the_report_lists_every_subject_in_the_disposition_table(tmp_path):
    res = _result_with_an_exclusion()
    text = _report_text(generate_report, res, tmp_path)
    for row in res["dispositions"]:
        assert str(row["subject"]) in text


def test_the_new_export_kinds_are_available_only_with_be_results():
    res = _result_with_an_exclusion()
    assert {"be", "be_disposition", "be_qc"} <= set(exporters.available(PharmState(be_results=res)))
    assert not {"be_disposition", "be_qc"} & set(exporters.available(PharmState()))


def test_the_disposition_csv_has_one_row_per_subject_and_parameter():
    res = _result_with_an_exclusion()
    rows = list(csv.DictReader(io.StringIO(exporters.export_csv(PharmState(be_results=res), "be_disposition"))))
    assert {"subject", "sequence_or_arm", "disposition", "parameter", "status", "reasons", "qc_flags"} <= set(rows[0])
    assert len(rows) == 13 * 3
    s3 = [r for r in rows if r["subject"] == "3" and r["parameter"] == "Cmax"][0]
    assert s3["status"] == "excluded" and "predose_gt_5pct_cmax" in s3["reasons"]


def test_the_qc_csv_lists_the_findings_with_their_basis():
    res = _result_with_an_exclusion()
    rows = list(csv.DictReader(io.StringIO(exporters.export_csv(PharmState(be_results=res), "be_qc"))))
    hit = [r for r in rows if r["rule"] == "predose_gt_5pct_cmax"]
    assert len(hit) == 1 and hit[0]["action"] == "excluded" and hit[0]["basis"] and hit[0]["subject"] == "3"


def test_the_be_csv_shows_each_parameters_role():
    res = _result_with_an_exclusion()
    rows = {r["parameter"]: r for r in csv.DictReader(io.StringIO(exporters.export_csv(
        PharmState(be_results=res), "be")))}
    assert [rows[k]["role"] for k in ("Cmax", "AUC_last", "AUC_inf")] == ["decision", "decision", "supportive"]


def test_a_formula_in_an_exclusion_reason_is_neutralised_in_the_csv():
    df = _frame(13)
    res = _call(df, **FULL, prespecified_exclusions=[{"subject": 4, "reason": "=cmd|' /C calc'!A0"}]).result
    text = exporters.export_csv(PharmState(be_results=res), "be_disposition")
    assert ",=cmd" not in text and "'=cmd" in text


def test_an_old_payload_without_roles_prints_no_empty_role_and_says_its_verdict_predates_them(tmp_path):
    res = _result_with_an_exclusion()
    old = {k: v for k, v in res.items() if k not in ("decision_parameters", "supportive_parameters", "decision_gap",
                                                     "qc_findings", "qc_summary", "dispositions")}
    old["parameters"] = {k: {kk: vv for kk, vv in v.items() if kk != "role"} for k, v in res["parameters"].items()}
    text = _report_text(generate_272, old, tmp_path)
    assert "()" not in text and "predates" in text
    assert "every reported parameter" not in text                          # no guess about what the old rule was
